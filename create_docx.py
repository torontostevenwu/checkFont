from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Add paragraph with Calibri font
p_calibri = doc.add_paragraph()
run_calibri = p_calibri.add_run("This is Calibri text.")
run_calibri.font.name = 'Calibri'
run_calibri.font.size = Pt(12)
p_calibri.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Add paragraph with Cambria font
p_cambria = doc.add_paragraph()
run_cambria = p_cambria.add_run("This is Cambria text.")
run_cambria.font.name = 'Cambria'
run_cambria.font.size = Pt(12)
p_cambria.alignment = WD_ALIGN_PARAGRAPH.LEFT

# Save the document
doc.save("sample.docx")

print("sample.docx created successfully.")
